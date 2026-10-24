"""
Tests for student upload template generation endpoint.
Complies with prompt B2 item 5 and agents.md §5, §7.
"""

import io

import openpyxl
from app.models.auth import User
from app.services.template import CANONICAL_COLUMNS
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
from tests.conftest import make_auth_headers


def get_headers(db: Session, role: str) -> dict[str, str]:
    user = db.query(User).filter(User.role == role).first()
    assert user is not None
    return make_auth_headers(user)


def test_template_unauthenticated(client: TestClient):
    resp = client.get("/v1/templates/student-upload.xlsx")
    assert resp.status_code == 401


def test_template_download_authenticated(client: TestClient, db_session: Session):
    headers = get_headers(db_session, "INSTITUTION_ADMIN")
    resp = client.get("/v1/templates/student-upload.xlsx", headers=headers)
    assert resp.status_code == 200
    assert (
        resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in resp.headers.get("content-disposition", "")
    assert "student-upload-template.xlsx" in resp.headers.get("content-disposition", "")

    # Load workbook into openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "StudentData" in wb.sheetnames
    assert "Instructions" in wb.sheetnames

    # Check StudentData headers
    ws_data = wb["StudentData"]
    headers = [
        ws_data.cell(row=1, column=col).value for col in range(1, len(CANONICAL_COLUMNS) + 1)
    ]
    expected_headers = [col[0] for col in CANONICAL_COLUMNS]
    assert headers == expected_headers

    # CRITICAL TEST: Check that APAAR column is strictly formatted as Text ('@')
    cell_a2 = ws_data.cell(row=2, column=1)
    cell_a10 = ws_data.cell(row=10, column=1)
    assert cell_a2.number_format == "@", f"Expected '@' (Text format), got {cell_a2.number_format}"
    assert cell_a10.number_format == "@", (
        f"Expected '@' (Text format), got {cell_a10.number_format}"
    )

    # Check Data Validation dropdown rules
    data_validations = ws_data.data_validations.dataValidation
    assert len(data_validations) > 0
    validation_formulas = [dv.formula1 for dv in data_validations if dv.formula1]
    assert any('"M,F,T"' in f for f in validation_formulas)
    assert any("GEN,EWS,OBC-NCL,SC,ST" in f for f in validation_formulas)

    # Check Instructions Sheet bilingual content
    ws_inst = wb["Instructions"]
    cell_title = ws_inst["A1"].value
    assert "Instructions" in cell_title
    assert "निर्देश" in cell_title

    # Check instruction rows have both English and Hindi text
    row5_en = ws_inst.cell(row=5, column=2).value
    row5_hi = ws_inst.cell(row=5, column=3).value
    assert "APAAR" in str(row5_en)
    assert "अपार" in str(row5_hi)
